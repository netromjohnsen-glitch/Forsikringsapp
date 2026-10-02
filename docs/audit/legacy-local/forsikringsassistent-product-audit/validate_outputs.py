import csv
import hashlib
import json
import pathlib
import zipfile
from collections import Counter

from docx import Document

BASE = pathlib.Path('/private/tmp/forsikringsassistent-product-audit')
full_path = BASE / 'product-comparison-audit.json'
root_path = BASE / 'root-cause-input.json'
matrix_path = BASE / 'product-comparison-matrix.csv'
findings_path = BASE / 'product-comparison-findings.csv'
md_path = BASE / 'product-comparison-audit.md'
docx_path = BASE / 'Forsikringsassistenten-produkt-audit.docx'

full = json.loads(full_path.read_text())
root = json.loads(root_path.read_text())
required = {'metadata','catalog','matrix','comparisons','findings','issue_signatures','aggregates','root_cause_clusters','pdf_customer_relevance','known_positive_controls','failures','limitations'}
assert required <= set(full), sorted(required - set(full))
assert full['aggregates']['planned_executions'] == 93
assert full['aggregates']['completed_executions'] == 93
assert full['aggregates']['failed_executions'] == 0
assert len(full['comparisons']) == 93
assert len(full['findings']) == 522
assert len(full['issue_signatures']) == 156
assert len(full['known_positive_controls']) == 7
assert full['failures'] == []
assert full['aggregates']['ai_calls'] == 0
assert full['aggregates']['pdf_operations'] == 0
assert full['aggregates']['runtime_web_requests'] == 0

with matrix_path.open(newline='') as f:
    matrix = list(csv.DictReader(f))
with findings_path.open(newline='') as f:
    findings = list(csv.DictReader(f))
assert len(matrix) == 93
assert len(findings) == 522
comparison_ids = {row['comparison_id'] for row in matrix}
assert len(comparison_ids) == 93
signature_ids = {item['issue_signature'] for item in full['issue_signatures']}
for row in findings:
    assert row['comparison_id'] in comparison_ids
    assert row['finding_id'] and row['reason_code'] and row['issue_signature']
    assert row['issue_signature'] in signature_ids
    assert row['raw_evidence_hash']
assert Counter(row['reason_code'] for row in findings) == Counter(full['aggregates']['reason_code_distribution'])

assert root_path.stat().st_size < full_path.stat().st_size
root_text = root_path.read_text()
assert 'normal_rows' not in root_text

md = md_path.read_text()
for marker in ['Executive summary','Audit matrix','Source and provenance','Limitations','93','522','156']:
    assert marker in md, marker
for forbidden in ['/Users/morten','/private/tmp/forsikringsassistent-product-audit']:
    assert forbidden not in md

assert zipfile.is_zipfile(docx_path)
doc = Document(docx_path)
docx_text = '\n'.join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        docx_text += '\n' + ' | '.join(cell.text for cell in row.cells)
for marker in ['Forsikringsassistenten – Produktsammenligning Audit','Executive summary','Audit matrix','Kilde og provenance','Begrensninger','ø','å']:
    assert marker in docx_text, marker
for forbidden in ['/Users/morten','/private/tmp/forsikringsassistent-product-audit']:
    assert forbidden not in docx_text

outputs = [full_path, root_path, matrix_path, findings_path, md_path, docx_path, BASE / 'audit-matrix-manifest.json', BASE / 'README.md']
manifest = {
    'audit_status': 'COMPLETE',
    'repository_commit': full['metadata'].get('git_head'),
    'counts': {
        'catalog_products': full['catalog']['total_products'],
        'eligible_products': full['catalog']['eligible_products'],
        'executions': len(matrix),
        'findings': len(findings),
        'issue_signatures': len(full['issue_signatures']),
    },
    'files': []
}
for path in outputs:
    data = path.read_bytes()
    manifest['files'].append({
        'file': path.name,
        'bytes': len(data),
        'sha256': hashlib.sha256(data).hexdigest(),
    })
(BASE / 'audit-output-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(manifest, ensure_ascii=False, indent=2))
