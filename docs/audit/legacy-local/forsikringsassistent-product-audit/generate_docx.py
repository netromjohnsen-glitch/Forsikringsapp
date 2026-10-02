import json, os, math
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/private/tmp/forsikringsassistent-product-audit')
AUDIT = json.loads((ROOT / 'product-comparison-audit.json').read_text('utf-8'))
OUT = ROOT / 'Forsikringsassistenten-produkt-audit.docx'
NAVY='17365D'; BLUE='DCE6F1'; PALE='F4F7FA'; MID='5B6573'; BORDER='D9D9D9'; BLACK='000000'; WHITE='FFFFFF'

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_border(cell, color=BORDER, size='4'):
    tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color)

def cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+name))
        if node is None: node=OxmlElement('w:'+name); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def keep_with_next(paragraph, value=True):
    pPr=paragraph._p.get_or_add_pPr(); node=pPr.find(qn('w:keepNext'))
    if value and node is None: node=OxmlElement('w:keepNext'); pPr.append(node)

def keep_together(paragraph):
    pPr=paragraph._p.get_or_add_pPr(); node=pPr.find(qn('w:keepLines'))
    if node is None: node=OxmlElement('w:keepLines'); pPr.append(node)

def set_col_width(cell, width):
    cell.width=Inches(width); tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(int(width*1440))); tcW.set(qn('w:type'),'dxa')

def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    hdr=table.rows[0]; set_repeat_header(hdr)
    for i,text in enumerate(headers):
        cell=hdr.cells[i]; set_cell_shading(cell,NAVY); set_cell_border(cell); cell_margins(cell)
        if widths: set_col_width(cell,widths[i])
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(str(text)); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font_size)
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cell=cells[i]; set_cell_border(cell); cell_margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: set_col_width(cell,widths[i])
            if ridx%2: set_cell_shading(cell,PALE)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
            r=p.add_run(str(value)); r.font.size=Pt(font_size); r.font.color.rgb=RGBColor.from_string(BLACK)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return table

def add_bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p

def add_label_paragraph(doc,label,text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(label+' '); r.bold=True; p.add_run(text); return p

def source_types(sources):
    values=[]
    for s in sources:
        label='IPID / produktark' if s.get('sourceType')=='ipid' else 'Produktside' if s.get('sourceType')=='product_page' else 'Offentlig vilkår'
        if label not in values: values.append(label)
    return ', '.join(values) or 'Ingen kilde vist'

doc=Document(); sec=doc.sections[0]
sec.page_height=Inches(11.69); sec.page_width=Inches(8.27); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string(BLACK)
styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.12
for name,size,space_before,space_after in [('Title',27,0,14),('Heading 1',18,18,8),('Heading 2',14,14,6),('Heading 3',11.5,10,4),('Heading 4',10,8,3)]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(BLACK); st.paragraph_format.space_before=Pt(space_before); st.paragraph_format.space_after=Pt(space_after); st.paragraph_format.keep_with_next=True
styles['Title'].paragraph_format.keep_with_next=True

# Header and footer
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.text='Forsikringsassistenten  Produkt audit'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MID)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run('Side '); r.font.size=Pt(8); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fp._p.append(field)

title=doc.add_paragraph(style='Title'); title.add_run('Forsikringsassistenten – Produktsammenligning Audit')
sub=doc.add_paragraph(); sub.add_run('Read only audit av katalogbasert produktsammenligning').bold=True
doc.add_paragraph(f"Auditdato {AUDIT['metadata']['audit_timestamp']}\nGit HEAD {AUDIT['metadata']['git_head']}\nArbeidskopi {AUDIT['metadata']['git_status_baseline']}")
doc.add_paragraph('Rapporten er et evidensgrunnlag for senere root cause analyse. Funnene er kontrollkandidater og fastslår ikke at forsikringssemantikken er feil.')

doc.add_heading('Executive summary',level=1)
agg=AUDIT['aggregates']; cat=AUDIT['catalog']
doc.add_paragraph(f"Auditen kjørte {agg['primary_comparisons']} primære sammenligninger og {agg['control_executions']} kontrollkjøringer på {cat['directly_exercised_products']} av {cat['eligible_products']} eligible katalogprodukter. Alle ni produktfamilier og alle eligible providers i hver familie ble representert. To historiske Eika Reise produkter ble korrekt utelatt.")
doc.add_paragraph(f"Detektoren registrerte {agg['flag_occurrences']} forekomster fordelt på {agg['unique_issue_signatures']} dedupliserte issue signatures. {agg['pdf_customer_relevance'].get('LIKELY_SHARED',0)} forekomster er klassifisert som mulige delte problemstillinger for produkt og PDF eller avtalesammenligning. Klassifiseringen er en arkitekturprioritering og ikke en bekreftet feil.")
add_table(doc,['Måling','Resultat'],[
    ['Katalogprodukter',cat['total_products']],['Eligible produkter',cat['eligible_products']],['Direkte testede produkter',cat['directly_exercised_products']],['Produktfamilier',len(cat['eligible_families'])],['Fullførte kjøringer',agg['completed_executions']],['Flaggforekomster',agg['flag_occurrences']],['Unike issue signatures',agg['unique_issue_signatures']],['Runtime nettverkskall / AI / PDF','0 / 0 / 0']
],[2.5,3.9],9)
doc.add_paragraph('Alle ti determinismekontroller og alle ni sidebytter var stabile. Ti representative kontroller rendret den faktiske React komponenten og samsvarte med runtime output.')

doc.add_heading('1 Overordnede funn',level=1)
add_table(doc,['Reason code','Forekomster'],sorted(agg['reason_code_distribution'].items(),key=lambda x:-x[1]),[4.7,1.7],8.5)
doc.add_heading('Tekniske klyngekandidater',level=2)
add_table(doc,['Lag','Signaturer','Forekomster','Familier'],[[c['layer'],c['unique_signatures'],c['occurrences'],', '.join(c['families'])] for c in sorted(AUDIT['root_cause_clusters'],key=lambda x:-x['occurrences'])],[2.0,.7,.8,2.9],7.8)
doc.add_paragraph('Klyngene viser hvor neste tekniske analyse bør begynne. HYPOTESE – IKKE VERIFISERT ROOT CAUSE.')
doc.add_heading('Ikke dokumentert analyse',level=2)
doc.add_paragraph(f"Runtime viste {AUDIT['unknown_analysis']['raw_occurrences']} unknown forekomster og {AUDIT['unknown_analysis']['unique_signatures']} unike side lokale signaturer.")
add_table(doc,['Klassifisering','Antall'],sorted(AUDIT['unknown_analysis']['classification_distribution'].items(),key=lambda x:-x[1]),[4.7,1.7],8.5)
doc.add_heading('Kjente positive kontroller',level=2)
add_table(doc,['Kontroll','Surfaced','Reason code'],[[c['name'],'JA' if c['surfaced'] else 'NEI',c['reason_code']] for c in AUDIT['known_positive_controls']],[3.7,.65,2.05],7.8)
review=AUDIT['metadata']['manual_sample_review']; doc.add_paragraph(f"Detektorkvalitet {review['detector_quality']}. Manuell kontroll omfattet {len(review['flagged_findings_reviewed'])} flaggede funn, {len(review['unflagged_rows_reviewed'])} uflaggede rader, {len(review['unknown_classifications_reviewed'])} unknown klassifiseringer, {len(review['provider_specific_candidates_reviewed'])} produktspesifikke kandidater og {len(review['provenance_chains_reviewed'])} provenance kjeder.")

family_numbers={'Bil':2,'Hus':3,'Innbo':4,'Reise':5,'Snøscooter':6,'Campingvogn':7,'Tilhenger':8,'MC':9,'Bobil':10}
for family in cat['eligible_families']:
    doc.add_page_break(); doc.add_heading(f"{family_numbers[family]} {family}",level=1)
    d=AUDIT['matrix']['by_family'][family]
    doc.add_paragraph(f"Providers testet {', '.join(d['tested_providers'])}. Eligible produkter {d['eligible_products']}. Primære sammenligninger {d['primary_comparisons']}. Kontrollkjøringer {d['control_executions']}. Flaggforekomster {d['flag_occurrences']}.")
    if d['common_reason_codes']: add_table(doc,['Reason code','Antall'],sorted(d['common_reason_codes'].items(),key=lambda x:-x[1]),[4.7,1.7],8)
    comparisons=[c for c in AUDIT['comparisons'] if c.get('status')=='COMPLETED' and c.get('primary') and c['insurance_type']==family]
    for c in comparisons:
        doc.add_heading(f"{c['first']['product']['company']} {c['first']['product']['product_name']} mot {c['second']['product']['company']} {c['second']['product']['product_name']}",level=2)
        add_label_paragraph(doc,'Comparison ID',c['comparison_id'])
        doc.add_paragraph(f"Seksjoner {c['metrics']['rendered_section_count']}. Rendererte fakta {c['metrics']['rendered_fact_count']}. Direkte rader {c['metrics']['directly_comparable_rows']}. Én sidige rader {c['metrics']['one_sided_rows']}. Unknown celler {c['metrics']['unknown_cells']}. Flagg {c['metrics']['flag_count']}.")
        fs_for=[f for f in AUDIT['findings'] if f['comparison_id']==c['comparison_id']]
        if not fs_for:
            doc.add_paragraph('Ingen auditflagg i denne sammenligningen.')
        else:
            add_table(doc,['Flagg','Seksjon og label','Side','Rå evidens'],[[f['reason_code'],f"{f['section']} / {f['visible_label']}",f['side'],f['raw_rendered_evidence']] for f in fs_for],[1.6,1.55,.55,2.7],7.1)

doc.add_page_break(); doc.add_heading('11 Tverrgående mønstre',level=1)
for c in sorted(AUDIT['root_cause_clusters'],key=lambda x:-x['occurrences']):
    doc.add_heading(c['layer'],level=2); doc.add_paragraph(f"{c['unique_signatures']} unike signaturer og {c['occurrences']} forekomster. Familier {', '.join(c['families'])}. PDF relevans {', '.join(f'{k} {v}' for k,v in c['pdf_relevance'].items())}. HYPOTESE – IKKE VERIFISERT ROOT CAUSE.")
doc.add_heading('Potensiell relevans for avtale og PDF sammenligning',level=2)
add_table(doc,['Klassifisering','Forekomster'],list(agg['pdf_customer_relevance'].items()),[4.7,1.7],8.5)
doc.add_paragraph('LIKELY_SHARED omfatter statusoppløsning, canonical identitet, parent child relasjoner, tilleggstilstand og mulig konseptoverlapp. PRODUCT_MODE_ONLY omfatter visningshierarki, interne etiketter, produktspesifikke detaljblokker og kundespesifikke referanser i ren produktmodus.')
doc.add_heading('Kun produktmodus og presentasjon',level=2)
doc.add_paragraph(f"{agg['reason_code_distribution'].get('DUPLICATE_DISPLAY_LABEL',0)} forekomster viser tre like nivåer i seksjon, gruppe og rad. {agg['reason_code_distribution'].get('INTERNAL_LABEL_LEAK',0)} forekomster viser interne avtaleetiketter. {agg['reason_code_distribution'].get('PROVIDER_SPECIFIC_NOT_COMPARABLE',0)} forekomster gjelder en sidige produktspesifikke detaljer.")

doc.add_page_break(); doc.add_heading('12 Kilde og provenance',level=1)
sp=AUDIT['source_provenance']; doc.add_paragraph(f"{sp['facts_checked']} fact source koblinger ble kontrollert. Det ble ikke funnet kandidatavvik for provider, forsikringstype, agreement scope, product ID, version eller manglende source record.")
add_table(doc,['Kildetype','Fact koblinger'],list(sp['source_type_distribution'].items()),[4.7,1.7],8.5)
doc.add_heading('Representativ provenance kontroll',level=2)
add_table(doc,['Familie','Produkt','Fact','Dokument','Punkt'],[[x['insurance_type'],f"{x['product']['company']} {x['product']['product_name']}",x['fact_key'],x['document_id'],f"{x.get('section','')} side {x.get('page','')}"] for x in review['provenance_chains_reviewed']],[.65,1.45,1.25,1.65,1.4],7.3)

doc.add_page_break(); doc.add_heading('13 Anbefalt input til neste root cause analyse',level=1)
doc.add_paragraph('Bruk root cause input JSON sammen med denne rapporten i en separat GPT 6 root cause og correctness oppgave. Neste oppgave bør kontrollere dedupliserte signaturer mot kode og katalogkilder før den bestemmer eventuelle rettelser og regresjonstester.')
for item in sorted(AUDIT['issue_signatures'],key=lambda x:(0 if x['potential_customer_pdf_relevance']=='LIKELY_SHARED' else 1,-x['occurrence_count']))[:15]:
    add_bullet(doc,f"{item['issue_signature']}  {item['reason_code']}  {item['occurrence_count']} forekomster  {item['required_next_investigation']}")
doc.add_paragraph('Ingen rettelser er implementert i denne auditen.')

doc.add_page_break(); doc.add_heading('14 Appendix og full comparison log',level=1)
doc.add_heading('Audit matrix',level=2)
add_table(doc,['Comparison ID','Familie','Produkt A','Produkt B','Type','Status'],[[c['comparison_id'],c['insurance_type'],f"{c.get('first',{}).get('product',{}).get('company','')} {c.get('first',{}).get('product',{}).get('product_name','')}",f"{c.get('second',{}).get('product',{}).get('company','')} {c.get('second',{}).get('product',{}).get('product_name','')}",'PRIMARY' if c.get('primary') else c.get('control_type',''),c['status']] for c in AUDIT['comparisons']],[2.25,.55,1.05,1.05,.8,.7],6.8)
doc.add_heading('Dedupliserte issue signatures med evidens',level=2)
for s in sorted(AUDIT['issue_signatures'],key=lambda x:(x['reason_code'],x['issue_signature'])):
    doc.add_heading(f"{s['reason_code']}  {s['issue_signature']}",level=3)
    doc.add_paragraph(f"Forekomster {s['occurrence_count']}. Familier {', '.join(s['families'])}. Providers {', '.join(s['providers'])}. Teknisk lag {s['likely_technical_layer']}. PDF relevans {s['potential_customer_pdf_relevance']}. Confidence {s['evidence_confidence']}.")
    add_label_paragraph(doc,'Comparison IDs',', '.join(s['comparison_ids']))
    for e in s['representative_evidence']:
        add_label_paragraph(doc,'Rå evidens',f"{e['section']} / {e['visible_label']} / {e['side']} / {e['raw_rendered_evidence']}")
        add_label_paragraph(doc,'Canonical og kilde',f"{e['canonical_key'] or 'ingen key'} / {e['source_document_ref'] or 'ingen dokumentreferanse i finding'}")
    add_label_paragraph(doc,'Neste kontroll',s['required_next_investigation'])

doc.add_heading('Begrensninger',level=2)
for item in AUDIT['limitations']: add_bullet(doc,item)
doc.add_paragraph('Auditen vurderer programvarerepresentasjon. Den sertifiserer ikke fullstendig juridisk dekningsfortolkning. De eligible katalogproduktene er ikke alle mulige norske forsikringsprodukter.')

doc.core_properties.title='Forsikringsassistenten Produktsammenligning Audit'
doc.core_properties.subject='Read only audit av katalogbasert produktsammenligning'
doc.core_properties.author='Forsikringsassistenten audit'
doc.save(OUT)
print(json.dumps({'status':'DOCX_READY','path':str(OUT),'bytes':OUT.stat().st_size},ensure_ascii=False))
